"""
The exam API holds question papers with their answer keys, so it belongs to
the people who set and run exams: teachers and school admins.

Every viewset here was gated on IsAuthenticated alone, so students and
parents were only kept out by whatever each get_queryset() happened to
return. Several actions skip get_queryset() entirely: they look records up by
an id from the request body with a bare Model.objects.get(id=...), which
reaches every school's rows.
"""

from datetime import date, time, timedelta
from textwrap import dedent

from django.contrib.auth import get_user_model
from django.test import SimpleTestCase
from rest_framework import status
from rest_framework.test import APITestCase

from academics.models import AcademicSession, EducationLevel, Term, TermType
from classroom.models import Class, GradeLevel
from exam.document_parser import ExamDocumentParser
from exam.models import (
    DifficultyLevel,
    Exam,
    ExamRegistration,
    ExamReview,
    ExamSchedule,
    ExamStatus,
    ExamType,
    QuestionBank,
)
from students.models import Student
from subject.models import Subject
from teacher.models import Teacher
from tenants.models import Tenant

User = get_user_model()

EXAMS = "/api/exams/exams/"


class ExamApiTestCase(APITestCase):
    def setUp(self):
        self.school = self.make_school("Exam School", "exam-school")
        self.other = self.make_school("Other Exam School", "other-exam-school")
        self.exam = self.make_exam(self.school)
        self.their_exam = self.make_exam(self.other)

    def make_school(self, name, slug):
        return Tenant.objects.create(
            name=name, slug=slug, status="active", is_active=True,
            owner_email=f"owner@{slug}.example.com")

    def grade_level(self, school):
        return GradeLevel.objects.filter(
            tenant=school, education_level__code="primary").order_by("order").first()

    def subject(self, school):
        subject, _ = Subject.objects.get_or_create(
            tenant=school, code="MATH-PRI",
            defaults={"name": "Mathematics", "education_levels": ["PRIMARY"]})
        return subject

    def make_exam(self, school, teacher=None):
        return Exam.objects.create(
            tenant=school, title="First Term Mathematics", subject=self.subject(school),
            grade_level=self.grade_level(school), teacher=teacher,
            exam_type=ExamType.objects.get(tenant=school, code="final_exam"),
            status=ExamStatus.objects.get(tenant=school, code="draft"),
            exam_date=date.today() + timedelta(days=7), start_time=time(9), end_time=time(11),
            objective_questions=[{
                "id": 1, "question": "2 + 2 = ?", "optionA": "3", "optionB": "4",
                "optionC": "5", "optionD": "6", "correctAnswer": "B", "marks": 1,
            }],
        )

    def make_user(self, role, school=None, **extra):
        school = school or self.school
        n = User.objects.count()
        user = User.objects.create_user(
            username=f"exam_{role}_{n}", email=f"exam_{role}_{n}@example.com", role=role,
            password="testpass123", is_active=True, tenant=school, **extra)
        if role == "teacher":
            Teacher.objects.get_or_create(
                user=user, defaults={"tenant": school, "employee_id": f"T{n}"})
        return user

    def login(self, role, school=None, **extra):
        user = self.make_user(role, school, **extra)
        self.client.force_authenticate(user=user)
        return user

    def make_student(self, school):
        level = EducationLevel.objects.get(tenant=school, code="primary")
        primary_1, _ = Class.objects.get_or_create(
            tenant=school, code="PRIMARY_1_EXAM",
            defaults={"name": "Primary 1 (exam)", "education_level": level,
                      "grade_number": 1, "order": 1})
        user = self.make_user("student", school)
        return Student.objects.create(
            user=user, gender="F", date_of_birth=date(2018, 1, 1),
            student_class=primary_1, tenant=school)

    def new_exam_payload(self, school):
        return {
            "title": "New exam", "subject": self.subject(school).id,
            "grade_level": self.grade_level(school).id,
            "exam_type": ExamType.objects.get(tenant=school, code="quiz").id,
            "status": ExamStatus.objects.get(tenant=school, code="draft").id,
            "exam_date": str(date.today() + timedelta(days=3)),
            "start_time": "09:00", "end_time": "10:00",
        }

    def get(self, url, school=None):
        return self.client.get(url, HTTP_X_TENANT_SLUG=(school or self.school).slug)

    def post(self, url, data=None, school=None):
        return self.client.post(
            url, data or {}, format="json", HTTP_X_TENANT_SLUG=(school or self.school).slug)


class FilteringTheExamListTest(ExamApiTestCase):
    """
    The screens filter by code ("quiz"), because a code means the same thing
    at every school while an id doesn't; other callers send ids. Both work,
    and a filter nobody can satisfy empties the list instead of refusing the
    whole request — sending a code to a filter that took only ids answered
    400 "Select a valid choice", so the exam list showed an error and no
    exams at all.
    """

    def setUp(self):
        super().setUp()
        self.login("admin")
        self.quiz = self.make_exam(self.school)
        self.quiz.title = "Pre-Nursery quiz"
        self.quiz.exam_type = ExamType.objects.get(tenant=self.school, code="quiz")
        self.quiz.status = ExamStatus.objects.get(tenant=self.school, code="pending_approval")
        self.quiz.save()

    def titles(self, query):
        response = self.get(f"{EXAMS}?{query}")
        self.assertEqual(response.status_code, status.HTTP_200_OK, response.data)
        results = response.data.get("results", response.data)
        return sorted(exam["title"] for exam in results)

    def test_a_type_or_status_can_be_named_by_its_code(self):
        self.assertEqual(self.titles("exam_type=quiz"), ["Pre-Nursery quiz"])
        self.assertEqual(self.titles("status=pending_approval"), ["Pre-Nursery quiz"])
        self.assertEqual(self.titles("exam_type=quiz&status=pending_approval"), ["Pre-Nursery quiz"])

    def test_an_id_still_works(self):
        quiz_type = ExamType.objects.get(tenant=self.school, code="quiz")
        pending = ExamStatus.objects.get(tenant=self.school, code="pending_approval")

        self.assertEqual(self.titles(f"exam_type={quiz_type.id}"), ["Pre-Nursery quiz"])
        self.assertEqual(self.titles(f"status={pending.id}"), ["Pre-Nursery quiz"])

    def test_a_type_nobody_has_finds_nothing_rather_than_refusing(self):
        self.assertEqual(self.titles("exam_type=no_such_type"), [])
        self.assertEqual(self.titles("status=no_such_status"), [])
        self.assertEqual(self.titles("difficulty_level=no_such_level"), [])

    def test_the_whole_list_comes_back_with_no_filter(self):
        self.assertEqual(self.titles(""), ["First Term Mathematics", "Pre-Nursery quiz"])

    def test_an_exam_set_online_is_in_the_list_like_any_other(self):
        """
        Asking for no filter asked, through an unticked checkbox, for exams
        that are not online and need no computer. Every exam set as either was
        missing from the list, with nothing on screen to explain it.
        """
        Exam.objects.filter(pk=self.quiz.pk).update(is_online=True, requires_computer=True)

        self.assertEqual(self.titles(""), ["First Term Mathematics", "Pre-Nursery quiz"])
        self.assertEqual(self.titles("exam_type=quiz"), ["Pre-Nursery quiz"])

    def test_online_can_still_be_asked_for_on_purpose(self):
        Exam.objects.filter(pk=self.quiz.pk).update(is_online=True)

        self.assertEqual(self.titles("is_online=true"), ["Pre-Nursery quiz"])
        self.assertEqual(self.titles("is_online=false"), ["First Term Mathematics"])


class SavingAnEditedExamTest(ExamApiTestCase):
    """
    The exam form reads an exam, changes a field and sends the lot back. Its
    type, status and difficulty come back as the nested objects the read
    serializer wrote, and its Status box works in codes. Only ids were
    accepted, so Save Changes answered 400 and nothing was saved — and the
    page showed the reason behind the open form, where nobody saw it.
    """

    def setUp(self):
        super().setUp()
        self.login("admin")

    def patch(self, data):
        return self.client.patch(f"{EXAMS}{self.exam.id}/", data, format="json",
                                 HTTP_X_TENANT_SLUG=self.school.slug)

    def test_the_status_can_come_back_as_the_object_it_was_read_as(self):
        approved = ExamStatus.objects.get(tenant=self.school, code="approved")

        response = self.patch({"status": {"id": approved.id, "code": "approved", "name": "Approved"}})

        self.assertEqual(response.status_code, status.HTTP_200_OK, response.data)
        self.exam.refresh_from_db()
        self.assertEqual(self.exam.status, approved)

    def test_a_code_names_a_status_or_a_type(self):
        response = self.patch({"status": "approved", "exam_type": "quiz"})

        self.assertEqual(response.status_code, status.HTTP_200_OK, response.data)
        self.exam.refresh_from_db()
        self.assertEqual((self.exam.status.code, self.exam.exam_type.code), ("approved", "quiz"))

    def test_an_id_still_works(self):
        quiz = ExamType.objects.get(tenant=self.school, code="quiz")

        response = self.patch({"exam_type": quiz.id})

        self.assertEqual(response.status_code, status.HTTP_200_OK, response.data)
        self.exam.refresh_from_db()
        self.assertEqual(self.exam.exam_type, quiz)

    def test_another_schools_status_is_refused(self):
        theirs = ExamStatus.objects.get(tenant=self.other, code="approved")

        response = self.patch({"status": theirs.id})

        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        self.exam.refresh_from_db()
        self.assertEqual(self.exam.status.tenant, self.school)

    def test_no_difficulty_chosen_is_saved_as_none(self):
        """The Difficulty box sits at 0 until one is picked, which is no row."""
        for empty in (0, "0", None):
            with self.subTest(sent=empty):
                response = self.patch({"difficulty_level": empty})

                self.assertEqual(response.status_code, status.HTTP_200_OK, response.data)
                self.exam.refresh_from_db()
                self.assertIsNone(self.exam.difficulty_level)

    def test_a_difficulty_can_still_be_set_by_code_or_id(self):
        hard = DifficultyLevel.objects.get(tenant=self.school, code="hard")

        by_code = self.patch({"difficulty_level": "hard"})
        self.exam.refresh_from_db()
        self.assertEqual((by_code.status_code, self.exam.difficulty_level), (status.HTTP_200_OK, hard))

        self.patch({"difficulty_level": None})
        by_id = self.patch({"difficulty_level": hard.id})
        self.exam.refresh_from_db()
        self.assertEqual((by_id.status_code, self.exam.difficulty_level), (status.HTTP_200_OK, hard))

    def test_a_status_nobody_has_is_refused_by_name(self):
        response = self.patch({"status": "no_such_status"})

        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertIn("status", response.data)


class ThePlainPageTest(ExamApiTestCase):
    """
    A school asked to type the paper on one page, the way they would in Word,
    breaking lines where they choose. The page is kept as typed, and so is the
    choice to print it in place of the question sections.
    """

    PAGE = "<p>Section A</p><p>1. What is 2 + 2?</p><p></p><p>A. 3</p><p>B. 4</p>"

    def setUp(self):
        super().setUp()
        self.login("admin")

    def test_a_new_exam_keeps_its_page_with_no_questions(self):
        response = self.post(EXAMS, {
            **self.new_exam_payload(self.school),
            "plain_page": self.PAGE, "print_plain_page": True,
        })

        self.assertEqual(response.status_code, status.HTTP_201_CREATED, response.data)
        exam = Exam.objects.get(tenant=self.school, title="New exam")
        self.assertEqual((exam.plain_page, exam.print_plain_page), (self.PAGE, True))
        self.assertEqual(exam.objective_questions, [])

    def test_the_page_comes_back_where_the_forms_and_printing_read_it(self):
        Exam.objects.filter(pk=self.exam.pk).update(plain_page=self.PAGE, print_plain_page=True)

        detail = self.get(f"{EXAMS}{self.exam.id}/").data
        listing = self.get(EXAMS).data
        listed = next(e for e in listing.get("results", listing) if e["id"] == self.exam.id)

        for read in (detail, listed):
            self.assertEqual((read["plain_page"], read["print_plain_page"]), (self.PAGE, True))

    def test_an_edit_can_change_the_page_and_stop_printing_it(self):
        Exam.objects.filter(pk=self.exam.pk).update(plain_page=self.PAGE, print_plain_page=True)

        response = self.client.patch(
            f"{EXAMS}{self.exam.id}/", {"plain_page": "<p>Changed</p>", "print_plain_page": False},
            format="json", HTTP_X_TENANT_SLUG=self.school.slug)

        self.assertEqual(response.status_code, status.HTTP_200_OK, response.data)
        self.exam.refresh_from_db()
        self.assertEqual((self.exam.plain_page, self.exam.print_plain_page), ("<p>Changed</p>", False))

    def test_an_exam_made_before_the_page_existed_has_none(self):
        self.assertEqual((self.exam.plain_page, self.exam.print_plain_page), ("", False))


class SavingAnExamsPrintSettingsTest(ExamApiTestCase):
    """
    The exam form and Print Preview both say an exam keeps its own print
    format, and both sent it - but the serializers left print_settings out,
    so it was dropped on every save and never read back. The paper always
    printed in the default format, whatever the school had chosen.
    """

    SETTINGS = {
        "font_family": "arial", "font_size": 11, "line_height": 1.15,
        "question_spacing": "compact", "option_layout": "stacked", "column_layout": 2,
        "margin": "custom", "margin_mm": {"top": 8, "right": 10, "bottom": 8, "left": 10},
        "show_marks": False, "show_instructions": True,
    }

    def setUp(self):
        super().setUp()
        self.login("admin")

    def patch(self, data):
        return self.client.patch(f"{EXAMS}{self.exam.id}/", data, format="json",
                                 HTTP_X_TENANT_SLUG=self.school.slug)

    def test_print_preview_saves_the_format_on_its_own(self):
        response = self.patch({"print_settings": self.SETTINGS})

        self.assertEqual(response.status_code, status.HTTP_200_OK, response.data)
        self.exam.refresh_from_db()
        self.assertEqual(self.exam.print_settings, self.SETTINGS)
        self.assertEqual(self.exam.title, "First Term Mathematics")

    def test_a_new_exam_keeps_its_format(self):
        response = self.post(EXAMS, {**self.new_exam_payload(self.school), "print_settings": self.SETTINGS})

        self.assertEqual(response.status_code, status.HTTP_201_CREATED, response.data)
        self.assertEqual(Exam.objects.get(tenant=self.school, title="New exam").print_settings, self.SETTINGS)

    def test_the_format_comes_back_where_the_form_and_printing_read_it(self):
        Exam.objects.filter(pk=self.exam.pk).update(print_settings=self.SETTINGS)

        detail = self.get(f"{EXAMS}{self.exam.id}/").data
        listing = self.get(EXAMS).data
        listed = next(e for e in listing.get("results", listing) if e["id"] == self.exam.id)

        self.assertEqual(detail["print_settings"], self.SETTINGS)
        self.assertEqual(listed["print_settings"], self.SETTINGS)

    def test_a_format_that_would_print_oddly_is_refused(self):
        for bad in (
            {"font_size": 40},
            {"font_family": "comic_sans"},
            {"line_height": 3},
            {"column_layout": 3},
            {"show_marks": "yes"},
            {"margin": "custom", "margin_mm": {"top": 1, "right": 10, "bottom": 10, "left": 10}},
            {"margin_mm": {"top": 10}},
            ["not", "an", "object"],
        ):
            with self.subTest(sent=bad):
                response = self.patch({"print_settings": bad})

                self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
                self.assertIn("print_settings", response.data)
        self.exam.refresh_from_db()
        self.assertEqual(self.exam.print_settings, {})


class ReadingATypedPaperTest(SimpleTestCase):
    """
    Pasted text and the Plain Page are both read by the same parser. With no
    title at the top, it took the first long line as the title and dropped it
    from the paper - and that line was "Section A" or the first question, so
    a whole section, or the first question, vanished from what was imported.
    """

    def sections(self, text):
        parsed = ExamDocumentParser(text.encode("utf-8"), "pasted-exam.txt").parse()
        return [(s["type"], len(s["questions"])) for s in parsed["sections"]]

    def test_a_paper_starting_with_section_a_keeps_section_a(self):
        self.assertEqual(self.sections(dedent("""\
            Section A - Objective

            What is 2 + 2?
            A. 2
            B. 3
            C. 4
            D. 5

            Section B - Theory
            1. Explain photosynthesis. (10 marks)
            """)), [("objective", 1), ("theory", 1)])

    def test_a_paper_starting_with_a_question_keeps_it(self):
        numbered = dedent("""\
            1. What is 2 + 2?
            A. 3
            B. 4
            C. 5
            D. 6
            2. What is 3 + 3?
            A. 6
            B. 7
            C. 8
            D. 9
            """)
        unnumbered = dedent("""\
            What is the capital of Nigeria?
            A. Lagos
            B. Abuja
            C. Kano
            D. Ibadan
            What is 3 + 3?
            A. 6
            B. 7
            C. 8
            D. 9
            """)

        self.assertEqual(self.sections(numbered), [("objective", 2)])
        self.assertEqual(self.sections(unnumbered), [("objective", 2)])

    def test_word_objective_questions_numbered_by_word_itself_are_kept(self):
        """
        A Word paper numbered its objective questions with Word's own list
        numbering, and put each question's options on line breaks inside the
        same paragraph. The "1." isn't in the text, so the reader found no
        questions and dropped Section A: only the theory section came in.
        """
        from io import BytesIO

        import docx

        paper = docx.Document()
        paper.add_heading("PRIMARY 4 ENGLISH STUDIES EXAMINATION", level=1)
        paper.add_heading("SECTION A: OBJECTIVE QUESTIONS", level=2)
        paper.add_paragraph("Choose the correct answer.")
        paper.add_paragraph("The boy _____ to school every morning.\nA. go\nB. goes\nC. going\nD. gone",
                            style="List Number")
        paper.add_paragraph("Choose the correct plural form of child.\nA. childs\nB. childes\nC. children\nD. childrens",
                            style="List Number")
        paper.add_heading("SECTION B: THEORY QUESTIONS", level=1)
        paper.add_paragraph("1. Write five common nouns you can find in your classroom.")
        file = BytesIO()
        paper.save(file)

        parsed = ExamDocumentParser(file.getvalue(), "paper.docx").parse()
        objective = parsed["sections"][0]

        self.assertEqual([(s["type"], len(s["questions"])) for s in parsed["sections"]],
                         [("objective", 2), ("theory", 1)])
        self.assertEqual(objective["instructions"], "Choose the correct answer.")
        self.assertEqual(objective["questions"][1]["question"], "<p>Choose the correct plural form of child.</p>")
        self.assertEqual(objective["questions"][1]["options"]["optionC"], "children")

    def word_answers(self, build):
        """The (options, answer) of each objective question in a Word paper."""
        from io import BytesIO

        import docx

        paper = docx.Document()
        paper.add_heading("SECTION A: OBJECTIVE QUESTIONS", level=2)
        build(paper)
        file = BytesIO()
        paper.save(file)
        parsed = ExamDocumentParser(file.getvalue(), "paper.docx").parse()
        return [q["correctAnswer"] for s in parsed["sections"] if s["type"] == "objective"
                for q in s["questions"]]

    @staticmethod
    def add_question(paper, stem, options, marked=None, how="bold"):
        from docx.enum.text import WD_COLOR_INDEX

        question = paper.add_paragraph(style="List Number")
        question.add_run(stem)
        for index, option in enumerate(options):
            question.add_run(f"\n{'ABCD'[index]}. ")
            run = question.add_run(option)
            if index == marked:
                if how == "bold":
                    run.bold = True
                elif how == "underline":
                    run.underline = True
                else:
                    run.font.highlight_color = WD_COLOR_INDEX.YELLOW

    def test_word_reads_a_bold_underlined_or_highlighted_option_as_the_answer(self):
        def build(paper):
            self.add_question(paper, "The boy _____ to school.", ["go", "goes", "going", "gone"], 1, "bold")
            self.add_question(paper, "Plural of child?", ["childs", "childes", "children", "childrens"], 2, "underline")
            self.add_question(paper, "Which is a noun?", ["quickly", "pretty", "teacher", "sing"], 2, "highlight")
            self.add_question(paper, "Nothing marked?", ["w", "x", "y", "z"])

        self.assertEqual(self.word_answers(build), ["B", "C", "C", None])

    def test_word_reads_a_marking_guide_and_leaves_its_theory_answers_alone(self):
        def build(paper):
            self.add_question(paper, "Question one?", ["w", "x", "y", "z"])
            self.add_question(paper, "Question two?", ["w", "x", "y", "z"])
            paper.add_heading("SECTION B: THEORY", level=2)
            paper.add_paragraph("1. Explain.")
            paper.add_heading("Marking Guide", level=2)
            for line in ("Section A", "1. C", "2-D", "Section B", "1. a. books b. boxes"):
                paper.add_paragraph(line)

        self.assertEqual(self.word_answers(build), ["C", "D"])

    def test_a_whole_question_in_bold_names_no_answer(self):
        """Every option bold is the teacher's styling, not one answer."""
        def build(paper):
            question = paper.add_paragraph(style="List Number")
            question.add_run("All bold?\nA. w\nB. x\nC. y\nD. z").bold = True

        self.assertEqual(self.word_answers(build), [None])

    def test_pasted_answer_lines_and_stars_name_the_answer(self):
        parsed = ExamDocumentParser(dedent("""\
            Section A - Objective
            What is 2 + 2?
            A. 3
            B. 4
            C. 5
            D. 6
            Answer: B
            What is 2 * 3?
            A. 5
            B. 6 *
            C. 7
            D. 8
            What is 1 + 1?
            A. 1
            B. 2
            C. 3
            D. 4
            Ans - (b)
            """).encode("utf-8"), "pasted-exam.txt").parse()
        questions = parsed["sections"][0]["questions"]

        self.assertEqual([q["correctAnswer"] for q in questions], ["B", "B", "B"])
        self.assertEqual(questions[1]["options"]["optionB"], "6")
        self.assertEqual(questions[1]["question"], "<p>What is 2 * 3?</p>")

    def test_a_title_line_is_still_left_out_of_the_questions(self):
        self.assertEqual(self.sections(dedent("""\
            First Term Mathematics Examination
            What is 2 + 2?
            A. 2
            B. 3
            C. 4
            D. 5
            """)), [("objective", 1)])


class ExamsAreForStaffTest(ExamApiTestCase):
    def test_students_and_parents_are_refused_every_exam_endpoint(self):
        reads = [
            EXAMS,
            f"{EXAMS}{self.exam.id}/",
            "/api/exams/schedules/",
            "/api/exams/registrations/",
            "/api/exams/results/",
            "/api/exams/statistics/",
            "/api/exams/exam-templates/",
            "/api/exams/exam-reviews/",
        ]
        writes = [
            EXAMS,
            "/api/exams/registrations/",
            "/api/exams/results/",
            "/api/exams/bulk-delete/",
            "/api/exams/bulk-update/",
            f"{EXAMS}bulk_create/",
            "/api/exams/registrations/bulk_register/",
            "/api/exams/registrations/mark_attendance/",
            "/api/exams/parse-text/",
            f"/api/exams/{self.exam.id}/export-pdf/",
        ]
        for role in ("student", "parent"):
            self.login(role)
            for url in reads:
                with self.subTest(role=role, method="GET", url=url):
                    response = self.get(url)
                    self.assertEqual(response.status_code, status.HTTP_403_FORBIDDEN)
                    self.assertNotIn(b"correctAnswer", response.content)
            for url in writes:
                with self.subTest(role=role, method="POST", url=url):
                    self.assertEqual(self.post(url).status_code, status.HTTP_403_FORBIDDEN)

    def test_students_cannot_create_exams(self):
        self.login("student")

        response = self.post(EXAMS, self.new_exam_payload(self.school))

        self.assertEqual(response.status_code, status.HTTP_403_FORBIDDEN)
        self.assertFalse(Exam.objects.filter(title="New exam").exists())

    def test_teachers_read_their_own_exams_with_answers(self):
        teacher = self.login("teacher")
        exam = self.make_exam(self.school, teacher=teacher.teacher)

        response = self.get(f"{EXAMS}{exam.id}/")

        self.assertEqual(response.status_code, status.HTTP_200_OK, response.data)
        self.assertEqual(response.data["objective_questions"][0]["correctAnswer"], "B")

    def test_school_admins_list_their_schools_exams(self):
        # Deliberately not is_staff: the role alone has to be enough.
        self.login("superadmin", is_staff=False)

        response = self.get(EXAMS)

        self.assertEqual(response.status_code, status.HTTP_200_OK, response.data)
        ids = [exam["id"] for exam in response.data["results"]]
        self.assertIn(self.exam.id, ids)
        self.assertNotIn(self.their_exam.id, ids)


class ExamActionsStayInTheirSchoolTest(ExamApiTestCase):
    def test_bulk_delete_does_not_reach_another_school(self):
        """The regression: any signed-in user could delete any school's exam by id."""
        self.login("superadmin")

        response = self.post("/api/exams/bulk-delete/", {"exam_ids": [self.their_exam.id, self.exam.id]})

        self.assertEqual(response.status_code, status.HTTP_200_OK, response.data)
        self.assertEqual(response.data["deleted_count"], 1)
        self.assertTrue(Exam.objects.filter(id=self.their_exam.id).exists())
        self.assertFalse(Exam.objects.filter(id=self.exam.id).exists())

    def test_bulk_update_does_not_reach_another_school(self):
        self.login("superadmin")

        response = self.post("/api/exams/bulk-update/", {
            "exam_ids": [self.their_exam.id, self.exam.id], "update_data": {"venue": "Hall B"}})

        self.assertEqual(response.status_code, status.HTTP_200_OK, response.data)
        self.assertEqual(response.data["updated_count"], 1)
        self.their_exam.refresh_from_db()
        self.exam.refresh_from_db()
        self.assertEqual(self.their_exam.venue, "")
        self.assertEqual(self.exam.venue, "Hall B")

    def test_bulk_update_cannot_move_an_exam_to_another_school(self):
        self.login("superadmin")

        response = self.post("/api/exams/bulk-update/", {
            "exam_ids": [self.exam.id], "update_data": {"tenant_id": str(self.other.id)}})

        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        self.exam.refresh_from_db()
        self.assertEqual(self.exam.tenant_id, self.school.id)

    def their_shared_question(self):
        their_teacher = self.make_user("teacher", self.other).teacher
        return QuestionBank.objects.create(
            tenant=self.other, created_by=their_teacher, question="Capital of Nigeria?",
            options=["Lagos", "Abuja", "Kano", "Ibadan"], correct_answer="Abuja",
            subject=self.subject(self.other), grade_level=self.grade_level(self.other),
            difficulty=DifficultyLevel.objects.get(tenant=self.other, code="easy"),
            is_shared=True)

    def test_teachers_cannot_import_another_schools_questions(self):
        """The regression: importing copied another school's answer key into this school's exam."""
        question = self.their_shared_question()
        self.login("teacher")

        response = self.post("/api/exams/question-bank/import_to_exam/", {
            "exam_id": self.exam.id, "question_ids": [question.id], "section_type": "objective"})

        self.assertEqual(response.status_code, status.HTTP_404_NOT_FOUND)
        self.exam.refresh_from_db()
        self.assertEqual(len(self.exam.objective_questions), 1)

    def test_teachers_cannot_import_into_another_schools_exam(self):
        teacher = self.login("teacher").teacher
        question = QuestionBank.objects.create(
            tenant=self.school, created_by=teacher, question="Capital of Nigeria?",
            options=["Lagos", "Abuja"], correct_answer="Abuja",
            subject=self.subject(self.school), grade_level=self.grade_level(self.school),
            difficulty=DifficultyLevel.objects.get(tenant=self.school, code="easy"))

        response = self.post("/api/exams/question-bank/import_to_exam/", {
            "exam_id": self.their_exam.id, "question_ids": [question.id], "section_type": "objective"})

        self.assertEqual(response.status_code, status.HTTP_404_NOT_FOUND)
        self.their_exam.refresh_from_db()
        self.assertEqual(len(self.their_exam.objective_questions), 1)

    def test_exams_cannot_be_created_from_another_schools_records(self):
        self.login("superadmin")
        payload = {**self.new_exam_payload(self.school), "subject": self.subject(self.other).id}

        response = self.post(EXAMS, payload)

        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertIn("subject", response.data)
        self.assertFalse(Exam.objects.filter(title="New exam").exists())

    def test_a_registration_cannot_point_at_another_schools_exam(self):
        """The regression: the new registration came back with that exam nested in it, answer key included."""
        ours = self.make_student(self.school)
        self.login("superadmin")

        response = self.post("/api/exams/registrations/", {"exam_id": self.their_exam.id, "student_id": ours.id})

        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertNotIn(b"correctAnswer", response.content)
        self.assertFalse(ExamRegistration.objects.exists())

    def test_registering_students_stays_in_the_school(self):
        ours = self.make_student(self.school)
        theirs = self.make_student(self.other)
        self.login("superadmin")

        refused = self.post("/api/exams/registrations/bulk_register/", {
            "exam_id": self.their_exam.id, "student_ids": [ours.id]})
        self.assertEqual(refused.status_code, status.HTTP_404_NOT_FOUND)

        one_by_one = self.post(f"{EXAMS}{self.exam.id}/register_student/", {"student_id": theirs.id})
        self.assertEqual(one_by_one.status_code, status.HTTP_404_NOT_FOUND)

        response = self.post("/api/exams/registrations/bulk_register/", {
            "exam_id": self.exam.id, "student_ids": [theirs.id, ours.id]})
        self.assertEqual(response.status_code, status.HTTP_200_OK, response.data)
        self.assertEqual(
            list(ExamRegistration.objects.values_list("exam", "student", "tenant")),
            [(self.exam.id, ours.id, self.school.id)])

    def test_mark_attendance_does_not_reach_another_school(self):
        theirs = ExamRegistration.objects.create(
            tenant=self.other, exam=self.their_exam, student=self.make_student(self.other))
        self.login("superadmin")

        response = self.post("/api/exams/registrations/mark_attendance/", {
            "attendance": [{"registration_id": theirs.id, "is_present": True}]})

        self.assertEqual(response.status_code, status.HTTP_200_OK, response.data)
        self.assertEqual(response.data["updated_count"], 0)
        theirs.refresh_from_db()
        self.assertFalse(theirs.is_present)

    def test_submit_for_review_does_not_reach_another_school(self):
        their_teacher = self.make_user("teacher", self.other).teacher
        self.login("teacher")

        response = self.post("/api/exams/exam-reviews/submit_for_review/", {
            "exam_id": self.their_exam.id, "reviewer_ids": [their_teacher.id]})

        self.assertEqual(response.status_code, status.HTTP_404_NOT_FOUND)
        self.assertFalse(ExamReview.objects.exists())

    def make_schedule(self, school, **extra):
        session = AcademicSession.objects.create(
            tenant=school, name="2026/2027", start_date=date(2026, 9, 7),
            end_date=date(2027, 7, 23), is_current=True)
        term_type, _ = TermType.objects.get_or_create(
            tenant=school, code="FT", defaults={"name": "First Term", "display_order": 1})
        term = Term.objects.create(
            tenant=school, term_type=term_type, academic_session=session,
            start_date=date(2026, 9, 7), end_date=date(2026, 12, 16), is_current=True)
        return ExamSchedule.objects.create(
            tenant=school, name="First Term Exams", academic_session=session, term=term,
            start_date=date(2026, 12, 1), end_date=date(2026, 12, 12), **extra)

    def test_a_default_schedule_does_not_clear_another_schools_default(self):
        theirs = self.make_schedule(self.other, is_default=True)
        ours = self.make_schedule(self.school)
        self.login("superadmin")

        response = self.post(f"/api/exams/schedules/{ours.id}/set_default/")

        self.assertEqual(response.status_code, status.HTTP_200_OK, response.data)
        ours.refresh_from_db()
        theirs.refresh_from_db()
        self.assertTrue(ours.is_default)
        self.assertTrue(theirs.is_default)
